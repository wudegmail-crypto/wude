"""公文格式处理工具 — 单文件版本"""
import customtkinter as ctk
from tkinter import filedialog, messagebox
import threading
import os
import json
import re
import datetime
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn

# ===================== 预设配置 =====================
PRESETS = {
    "gbt_gov": {
        "name": "GB/T公文标准", "font_cn": "仿宋_GB2312", "font_en": "Times New Roman",
        "font_size_cn": 16, "font_size_en": 16, "line_spacing": 28.0, "line_spacing_rule": "FIXED",
        "margin_top": 3.7, "margin_bottom": 3.5, "margin_left": 2.8, "margin_right": 2.6,
        "first_line_indent_chars": 2, "title_font_size": 22, "title_bold": True, "title_align": "center",
        "heading1_font_size": 16, "heading1_bold": True, "heading2_font_size": 16, "heading2_bold": False,
    },
    "academic": {
        "name": "学术论文", "font_cn": "宋体", "font_en": "Times New Roman",
        "font_size_cn": 12, "font_size_en": 12, "line_spacing": 1.5, "line_spacing_rule": "MULTIPLE",
        "margin_top": 2.54, "margin_bottom": 2.54, "margin_left": 3.17, "margin_right": 3.17,
        "first_line_indent_chars": 2, "title_font_size": 16, "title_bold": True, "title_align": "center",
        "heading1_font_size": 14, "heading1_bold": True, "heading2_font_size": 12, "heading2_bold": True,
    },
    "legal": {
        "name": "法律文书", "font_cn": "仿宋", "font_en": "Times New Roman",
        "font_size_cn": 14, "font_size_en": 14, "line_spacing": 25.0, "line_spacing_rule": "FIXED",
        "margin_top": 2.5, "margin_bottom": 2.5, "margin_left": 2.5, "margin_right": 2.5,
        "first_line_indent_chars": 2, "title_font_size": 18, "title_bold": True, "title_align": "center",
        "heading1_font_size": 14, "heading1_bold": True, "heading2_font_size": 14, "heading2_bold": False,
    },
    "custom": {
        "name": "自定义", "font_cn": "宋体", "font_en": "Times New Roman",
        "font_size_cn": 12, "font_size_en": 12, "line_spacing": 1.5, "line_spacing_rule": "MULTIPLE",
        "margin_top": 2.54, "margin_bottom": 2.54, "margin_left": 3.17, "margin_right": 3.17,
        "first_line_indent_chars": 2, "title_font_size": 16, "title_bold": True, "title_align": "center",
        "heading1_font_size": 14, "heading1_bold": True, "heading2_font_size": 12, "heading2_bold": True,
    },
}

# ===================== 标点修复 =====================
CN_PUNCT = {',': '，', '.': '。', ';': '；', ':': '：', '?': '？', '!': '！',
            '(': '（', ')': '）', '<': '《', '>': '》', '[': '【', ']': '】'}


def _is_chinese(ch):
    return '一' <= ch <= '鿿'


def fix_punctuation(text):
    """修复单段文本标点"""
    chars = list(text)
    for i, ch in enumerate(chars):
        if ch in CN_PUNCT:
            nearby_cn = (i > 0 and _is_chinese(text[i - 1])) or \
                       (i < len(text) - 1 and _is_chinese(text[i + 1]))
            if nearby_cn and not (ch in '.,' and i > 0 and text[i - 1].isdigit()):
                chars[i] = CN_PUNCT[ch]

    # 引号配对
    result = ''.join(chars)
    # 中文数字后的逗号转为枚举顿号
    result = re.sub(r'([一二三四五六七八九十])([，,])', r'\1、', result)
    quote_open = True
    chars2 = list(result)
    for i, ch in enumerate(chars2):
        if ch == '"':
            if i > 0 and _is_chinese(result[i - 1]) or i < len(result) - 1 and _is_chinese(result[i + 1]):
                chars2[i] = '“' if quote_open else '”'
                quote_open = not quote_open
    result = ''.join(chars2)
    return result



def fix_doc_punctuation(doc):
    """修复全文标点，返回修复处数"""
    count = 0
    for para in doc.paragraphs:
        for run in para.runs:
            old = run.text
            new = fix_punctuation(old)
            if old != new:
                run.text = new
                count += 1
    return count


# ===================== 格式诊断 =====================
def check_document(doc, preset):
    """诊断文档格式问题"""
    issues = []
    for i, section in enumerate(doc.sections):
        checks = [
            ('上边距', section.top_margin, preset['margin_top']),
            ('下边距', section.bottom_margin, preset['margin_bottom']),
            ('左边距', section.left_margin, preset['margin_left']),
            ('右边距', section.right_margin, preset['margin_right']),
        ]
        for name, actual, expected in checks:
            actual_cm = round(actual / 360000, 2) if actual else 0
            ok = abs(actual_cm - expected) <= 0.1
            issues.append(f"{'✅' if ok else '⚠'} 第{i+1}节{name}: {actual_cm}cm {'(应为'+str(expected)+'cm)' if not ok else ''}")

    total_paras = len(doc.paragraphs)
    prev_empty = False
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        n = i + 1

        if not text:
            if prev_empty:
                issues.append(f"⚠ 第{n}段: 多余空行（连续空段）")
            prev_empty = True
            continue
        prev_empty = False
        has_cn = any(_is_chinese(c) for c in text)
        pf = para.paragraph_format

        is_doc_title = _is_doc_title(text, i, total_paras)
        is_h1 = _is_heading1(text)
        is_h2 = _is_heading2(text)

        if has_cn:
            expected_font = preset['font_cn']
            for run in para.runs:
                if run.font.name and run.font.name != expected_font:
                    role = '（标题）' if is_doc_title else '（一级标题）' if is_h1 else ''
                    issues.append(f"⚠ 第{n}段{role}: 字体'{run.font.name}', 应为'{expected_font}'")
                    break

        if is_doc_title:
            expected_size = preset.get('title_font_size', 22)
        elif is_h1 or is_h2:
            expected_size = preset.get('heading1_font_size', preset['font_size_cn'])
        else:
            expected_size = preset['font_size_cn']

        for run in para.runs:
            if run.font.size:
                actual_pt = round(run.font.size / 12700, 1)
                if abs(actual_pt - expected_size) > 0.5:
                    role = '（标题）' if is_doc_title else '（一级标题）' if is_h1 else ''
                    issues.append(f"⚠ 第{n}段{role}: 字号{actual_pt}pt, 应为{expected_size}pt")
                    break

        if pf.line_spacing:
            if pf.line_spacing_rule == WD_LINE_SPACING.EXACTLY:
                actual_ls = round(pf.line_spacing / 12700, 1)
                unit = '磅'
            else:
                actual_ls = round(pf.line_spacing, 1) if isinstance(pf.line_spacing, float) else pf.line_spacing
                unit = '倍'
            expected_ls = preset['line_spacing']
            expected_unit = '磅' if preset['line_spacing_rule'] == 'FIXED' else '倍'
            if abs(actual_ls - expected_ls) > 0.5:
                issues.append(f"⚠ 第{n}段: 行距{actual_ls}{unit}, 应为{expected_ls}{expected_unit}")

        if pf.line_spacing is None and pf.line_spacing_rule is None:
            expected_ls = preset['line_spacing']
            unit = '磅' if preset['line_spacing_rule'] == 'FIXED' else '倍'
            issues.append(f"⚠ 第{n}段: 行距未设置（默认单倍）, 应为{expected_ls}{unit}")

        if not is_doc_title and not is_h1 and not is_h2:
            if re.match(r'^[一-鿿]+[,，：\s]', text) and len(text) < 30:
                pass
            else:
                expected_indent = preset['font_size_cn'] * preset.get('first_line_indent_chars', 2)
                actual_indent = pf.first_line_indent
                if actual_indent is not None:
                    actual_pt = round(actual_indent / 12700, 1)
                    if abs(actual_pt - expected_indent) > 2:
                        issues.append(f"⚠ 第{n}段: 首行缩进{actual_pt}pt, 应为{expected_indent}pt（{preset.get('first_line_indent_chars', 2)}字符）")
                else:
                    issues.append(f"⚠ 第{n}段: 首行缩进未设置, 应为{expected_indent}pt")

        if is_doc_title:
            title_should_bold = preset.get('title_bold', True)
            for run in para.runs:
                if run.bold and not title_should_bold:
                    issues.append(f"⚠ 第{n}段（标题）: 不应加粗（GB/T标准）")
                    break

        if is_h1:
            h1_should_bold = preset.get('heading1_bold', True)
            if h1_should_bold and not any(r.bold for r in para.runs if r.text.strip()):
                issues.append(f"⚠ 第{n}段（一级标题）: 应加粗（黑体）")

        en_puncts = [',', '.', ';', ':', '?', '!']
        found = [p for p in en_puncts if text.count(p) > 0]
        if found:
            issues.append(f"⚠ 第{n}段: 英文标点混用 {', '.join(f'{p}×{text.count(p)}' for p in found)}")

    # 落款对齐检查（在尾部段落中判断）
    for i, para in enumerate(doc.paragraphs):
        if i < total_paras - 3:
            continue
        text = para.text.strip()
        if not text:
            continue
        n = i + 1
        if re.match(r'^[A-Za-z一-鿿]{2,}(公司|集团|中心|局|部|委|办|处|院|所)', text):
            if para.alignment != WD_ALIGN_PARAGRAPH.RIGHT:
                issues.append(f"⚠ 第{n}段: 发文机关署名应右对齐")
        if re.match(r'^\d{4}年\d{1,2}月\d{1,2}日', text):
            if para.alignment != WD_ALIGN_PARAGRAPH.RIGHT:
                issues.append(f"⚠ 第{n}段: 成文日期应右对齐")

    return issues


# ===================== 格式应用 =====================
def _set_font(run, cn_font, en_font, size):
    """设置字体（含东亚字体）"""
    run.font.size = Pt(size)
    run.font.name = cn_font
    r = run._element
    rPr = r.find(qn('w:rPr'))
    if rPr is None:
        rPr = r.makeelement(qn('w:rPr'), {})
        r.insert(0, rPr)
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = rPr.makeelement(qn('w:rFonts'), {})
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:eastAsia'), cn_font)
    rFonts.set(qn('w:ascii'), en_font)
    rFonts.set(qn('w:hAnsi'), en_font)


def _is_heading1(text):
    """一级标题：一、二、三、"""
    return bool(re.match(r'^[一二三四五六七八九十]+、', text))


def _is_heading2(text):
    """二级标题：（一）（二）"""
    return bool(re.match(r'^（[一二三四五六七八九十]+）', text))


def _is_doc_title(text, para_index, total_paras):
    """公文主标题：文档前部的关于...的通知/函/报告等"""
    if not text or len(text) > 80 or para_index > 5:
        return False
    return bool(re.search(r'关于.*的(通知|函|报告|请示|批复|决定|意见|通报|纪要|公告|通告|议案)', text))


def _is_title(text, preset):
    """判断是否为标题段落（用于格式应用时的通用检测）"""
    if not text or len(text) > 80:
        return False
    patterns = [r'^第[一二三四五六七八九十\d]+[章节部分]', r'^[一二三四五六七八九十]、',
                r'^[\d]+\.[\d]', r'^（[一二三四五六七八九十\d]）', r'^[1-9]、']
    return any(re.match(p, text) for p in patterns)


def _is_main_recipient(text):
    """主送机关：各省，自治区，直辖市分公司："""
    return bool(re.match(r'^[一-鿿A-Za-z]+[，,：:][\s\S]*[：:]$', text) and len(text) < 40)


def _is_signature(text, para_index, total):
    """落款：发文机关署名或成文日期"""
    if para_index < total - 5:
        return False
    if re.match(r'^[A-Za-z一-鿿]{2,}(公司|集团|中心|局|部|委|办|处|院|所)', text):
        return True
    if re.match(r'^\d{4}年\d{1,2}月\d{1,2}日', text):
        return True
    return False


def apply_format(doc, preset):
    """应用预设格式"""
    para_count = title_count = heading_count = sig_count = 0
    total = len(doc.paragraphs)

    for section in doc.sections:
        section.top_margin = Cm(preset['margin_top'])
        section.bottom_margin = Cm(preset['margin_bottom'])
        section.left_margin = Cm(preset['margin_left'])
        section.right_margin = Cm(preset['margin_right'])

    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if not text:
            continue
        para_count += 1
        pf = para.paragraph_format

        if preset['line_spacing_rule'] == 'FIXED':
            pf.line_spacing = Pt(preset['line_spacing'])
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        else:
            pf.line_spacing = preset['line_spacing']

        is_doc_title = _is_doc_title(text, i, total)
        is_h1 = _is_heading1(text)
        is_h2 = _is_heading2(text)
        is_sig = _is_signature(text, i, total)
        is_recipient = _is_main_recipient(text)

        if is_sig:
            sig_count += 1
            pf.first_line_indent = Pt(0)
            pf.space_before = Pt(6)
            pf.space_after = Pt(0)
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for run in para.runs:
                run.bold = False
                _set_font(run, preset['font_cn'], preset['font_en'], preset['font_size_cn'])

        elif is_doc_title:
            title_count += 1
            pf.first_line_indent = Pt(0)
            pf.space_before = Pt(12)
            pf.space_after = Pt(12)
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                run.bold = preset.get('title_bold', True)
                _set_font(run, preset['font_cn'], preset['font_en'], preset['title_font_size'])

        elif is_h1:
            heading_count += 1
            pf.first_line_indent = Pt(0)
            pf.space_before = Pt(12)
            pf.space_after = Pt(6)
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in para.runs:
                run.bold = preset.get('heading1_bold', True)
                _set_font(run, preset['font_cn'], preset['font_en'],
                         preset.get('heading1_font_size', preset['font_size_cn']))

        elif is_h2:
            heading_count += 1
            pf.first_line_indent = Pt(0)
            pf.space_before = Pt(6)
            pf.space_after = Pt(3)
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in para.runs:
                run.bold = preset.get('heading2_bold', False)
                _set_font(run, preset['font_cn'], preset['font_en'],
                         preset.get('heading2_font_size', preset['font_size_cn']))

        elif is_recipient:
            pf.first_line_indent = Pt(0)
            pf.space_before = Pt(0)
            pf.space_after = Pt(6)
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in para.runs:
                run.bold = False
                _set_font(run, preset['font_cn'], preset['font_en'], preset['font_size_cn'])

        else:
            pf.first_line_indent = Pt(preset['font_size_cn']) * preset['first_line_indent_chars']
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            for run in para.runs:
                _set_font(run, preset['font_cn'], preset['font_en'], preset['font_size_cn'])

    return para_count, title_count, heading_count, sig_count


# ===================== 处理引擎 =====================
def process(input_path, output_path, mode, preset_key):
    """主处理流程"""
    preset = PRESETS[preset_key]
    log = []
    t0 = datetime.datetime.now()
    mode_names = {'full': '智能一键处理', 'diagnose': '格式诊断', 'punctuation': '标点修复'}

    log.append(f"[{t0:%H:%M:%S}] 开始处理: {os.path.basename(input_path)}")
    log.append(f"[{t0:%H:%M:%S}] 预设: {preset['name']} | 模式: {mode_names.get(mode, mode)}")

    try:
        doc = Document(input_path)
    except Exception as e:
        log.append(f"❌ 无法打开文件: {e}")
        return log

    if mode in ('full', 'punctuation'):
        n = fix_doc_punctuation(doc)
        log.append(f"[{datetime.datetime.now():%H:%M:%S}] 标点修复: {n}处")

    if mode == 'diagnose':
        issues = check_document(doc, preset)
        log.append(f"[{datetime.datetime.now():%H:%M:%S}] 诊断完成: {len(issues)}项")
        log.extend(issues)
        return log

    if mode == 'full':
        paras, titles, headings, sigs = apply_format(doc, preset)
        log.append(f"[{datetime.datetime.now():%H:%M:%S}] 格式应用: 段落{paras}个, 主标题{titles}个, 层级标题{headings}个, 落款{sigs}个")

    try:
        doc.save(output_path)
        elapsed = (datetime.datetime.now() - t0).total_seconds()
        log.append(f"✅ 处理完成 → {output_path} ({elapsed:.1f}s)")
    except Exception as e:
        log.append(f"❌ 保存失败: {e}")

    return log


# ===================== GUI =====================
class App(ctk.CTk):
    def __init__(self):
        super().__init__()
        ctk.set_appearance_mode("dark")
        ctk.set_default_color_theme("blue")

        self.title("公文格式处理工具 v1.0")
        self.geometry("760x820")
        self.minsize(680, 700)

        self.preset_key = "gbt_gov"
        self.mode = ctk.StringVar(value="智能一键处理")
        self.input_path = ctk.StringVar()
        self.output_path = ctk.StringVar()
        self._build()
        self._update_params()

    def _build(self):
        # 标题
        ctk.CTkLabel(self, text="📄 公文格式处理工具", font=ctk.CTkFont(size=22, weight="bold")).pack(pady=(20, 5))
        ctk.CTkLabel(self, text="支持 .docx 文档 | 智能标点修复 | 标准格式一键应用",
                    text_color="#888").pack(pady=(0, 18))

        # 文件选择
        f1 = ctk.CTkFrame(self)
        f1.pack(fill="x", padx=20, pady=(0, 10))
        ctk.CTkLabel(f1, text="📂 文件设置", font=ctk.CTkFont(size=13, weight="bold")).pack(anchor="w", padx=15, pady=(10, 3))
        for label, var, btn_text, is_input in [
            ("输入文件", self.input_path, "选择文件", True),
            ("输出位置", self.output_path, "选择目录", False)]:
            row = ctk.CTkFrame(f1, fg_color="transparent")
            row.pack(fill="x", padx=15, pady=3)
            ctk.CTkLabel(row, text=label, width=65).pack(side="left")
            ctk.CTkEntry(row, textvariable=var, height=32).pack(side="left", fill="x", expand=True, padx=(8, 8))
            ctk.CTkButton(row, text=btn_text, width=85, height=32,
                         command=lambda i=is_input: self._browse(i)).pack(side="right")

        # 处理模式
        f2 = ctk.CTkFrame(self)
        f2.pack(fill="x", padx=20, pady=(0, 10))
        ctk.CTkLabel(f2, text="⚙ 处理模式", font=ctk.CTkFont(size=13, weight="bold")).pack(anchor="w", padx=15, pady=(10, 3))
        ctk.CTkSegmentedButton(f2, values=["智能一键处理", "格式诊断", "标点修复"],
                               variable=self.mode, height=34, command=self._on_mode_change).pack(fill="x", padx=15, pady=(5, 5))
        self.mode_desc = ctk.CTkLabel(f2, text="", text_color="#aaa", font=ctk.CTkFont(size=11),
                                       wraplength=700, justify="left")
        self.mode_desc.pack(anchor="w", padx=18, pady=(0, 8))
        self._on_mode_change("智能一键处理")

        # 格式预设
        f3 = ctk.CTkFrame(self)
        f3.pack(fill="x", padx=20, pady=(0, 10))
        ctk.CTkLabel(f3, text="📐 格式预设", font=ctk.CTkFont(size=13, weight="bold")).pack(anchor="w", padx=15, pady=(10, 3))
        preset_keys = list(PRESETS.keys())
        preset_names = [PRESETS[k]["name"] for k in preset_keys]
        preset_seg = ctk.CTkSegmentedButton(f3, values=preset_names,
                                            command=self._on_preset_seg, height=32)
        preset_seg.pack(fill="x", padx=15, pady=(5, 3))
        preset_seg.set(PRESETS["gbt_gov"]["name"])
        self.preset_desc = ctk.CTkLabel(f3, text="", text_color="#aaa", font=ctk.CTkFont(size=11),
                                         wraplength=700, justify="left")
        self.preset_desc.pack(anchor="w", padx=18, pady=(0, 8))
        self._update_preset_desc()

        # 参数网格
        self.param_vars = {}
        self.param_entries = {}
        pg = ctk.CTkFrame(f3, fg_color="transparent")
        pg.pack(fill="x", padx=15, pady=(5, 10))
        fields = [("字体", "font_cn"), ("字号", "font_size_cn"), ("行距", "line_spacing"),
                  ("缩进(字符)", "first_line_indent_chars"), ("上边距(cm)", "margin_top"),
                  ("下边距(cm)", "margin_bottom"), ("左边距(cm)", "margin_left"), ("右边距(cm)", "margin_right")]
        for idx, (label, key) in enumerate(fields):
            r, c = idx // 4, idx % 4
            cell = ctk.CTkFrame(pg, fg_color="transparent")
            cell.grid(row=r, column=c, sticky="w", padx=5, pady=3)
            ctk.CTkLabel(cell, text=label, font=ctk.CTkFont(size=11)).pack(side="left", padx=(0, 4))
            var = ctk.StringVar()
            ent = ctk.CTkEntry(cell, textvariable=var, width=75, height=28, state="readonly")
            ent.pack(side="left")
            self.param_vars[key] = var
            self.param_entries[key] = ent

        # 开始按钮
        self.btn = ctk.CTkButton(self, text="开始处理", height=40,
                                 font=ctk.CTkFont(size=14, weight="bold"),
                                 command=self._start)
        self.btn.pack(pady=(5, 10), padx=20)

        # 日志
        f4 = ctk.CTkFrame(self)
        f4.pack(fill="both", expand=True, padx=20, pady=(0, 8))
        ctk.CTkLabel(f4, text="处理日志", font=ctk.CTkFont(weight="bold")).pack(anchor="w", padx=15, pady=(10, 3))
        self.log = ctk.CTkTextbox(f4, wrap="word", font=ctk.CTkFont(family="Consolas", size=11))
        self.log.pack(fill="both", expand=True, padx=10, pady=(0, 10))
        self.log.configure(state="disabled")

        # 状态栏
        self.status = ctk.CTkLabel(self, text="就绪", anchor="w", height=28,
                                   fg_color="#2b2b2b", corner_radius=0, padx=15)
        self.status.pack(fill="x", side="bottom")

    def _browse(self, is_input):
        if is_input:
            p = filedialog.askopenfilename(filetypes=[("Word文档", "*.docx"), ("所有文件", "*.*")])
            if p:
                self.input_path.set(p)
                base = os.path.splitext(os.path.basename(p))[0]
                self.output_path.set(os.path.join(os.path.dirname(p), f"{base}_已处理.docx"))
        else:
            p = filedialog.askdirectory()
            if p:
                self.output_path.set(p)

    def _on_mode_change(self, name):
        descs = {
            "智能一键处理": "标点修复 → 格式应用：自动修复中英文标点混用，并应用下方选择的格式标准，一步到位完成文档规范化处理。",
            "格式诊断": "仅检查不修改：分析文档的字体、字号、行距、页边距、标点等是否符合标准，输出诊断报告，不修改原文件。",
            "标点修复": "仅修复标点：将中文上下文中的英文逗号、句号、引号等替换为中文标点符号，不改变文档格式。",
        }
        self.mode_desc.configure(text=descs.get(name, ""))

    def _on_preset_seg(self, name):
        for k, v in PRESETS.items():
            if v["name"] == name:
                self.preset_key = k
                self._update_params()
                self._update_preset_desc()
                return

    def _update_preset_desc(self):
        descs = {
            "gbt_gov": "正文：仿宋_GB2312 三号，行距28磅，页边距上3.7/下3.5/左2.8/右2.6cm，首行缩进2字符，标题22磅居中加粗。",
            "academic": "正文：宋体 小四，1.5倍行距，页边距上下2.54/左右3.17cm，首行缩进2字符，标题16磅居中加粗。",
            "legal": "正文：仿宋 四号，固定行距25磅，页边距各2.5cm，首行缩进2字符，标题18磅居中加粗。",
            "custom": '自定义参数：手动调整下方各项格式参数，调整后点击"开始处理"即可应用。',
        }
        self.preset_desc.configure(text=descs.get(self.preset_key, ""))

    def _update_params(self):
        cfg = PRESETS[self.preset_key]
        for k in self.param_vars:
            self.param_vars[k].set(str(cfg.get(k, "")))
            self.param_entries[k].configure(state="normal" if self.preset_key == "custom" else "readonly")

    def _start(self):
        inp = self.input_path.get().strip()
        out = self.output_path.get().strip()
        if not inp:
            return messagebox.showwarning("提示", "请选择输入文件")
        if not os.path.exists(inp):
            return messagebox.showerror("错误", "文件不存在")
        mode_map = {"智能一键处理": "full", "格式诊断": "diagnose", "标点修复": "punctuation"}
        mode = mode_map.get(self.mode.get(), "full")
        if mode != "diagnose" and not out:
            return messagebox.showwarning("提示", "请选择输出位置")

        if self.preset_key == "custom":
            for k, v in self.param_vars.items():
                try:
                    PRESETS["custom"][k] = float(v.get()) if '.' in v.get() else int(v.get())
                except ValueError:
                    PRESETS["custom"][k] = v.get()

        self.btn.configure(state="disabled")
        self.status.configure(text="处理中...")
        self.log.configure(state="normal")
        self.log.delete("1.0", "end")
        self.log.configure(state="disabled")

        def run():
            logs = process(inp, out, mode, self.preset_key)
            self.after(0, lambda: self._done(logs))

        threading.Thread(target=run, daemon=True).start()

    def _done(self, logs):
        self.log.configure(state="normal")
        for line in logs:
            self.log.insert("end", line + "\n")
        self.log.see("end")
        self.log.configure(state="disabled")
        self.status.configure(text="处理完成" if any("✅" in l for l in logs) else "处理失败")
        self.btn.configure(state="normal")


if __name__ == "__main__":
    App().mainloop()
