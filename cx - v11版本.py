import re
import os
import shutil
import traceback
from copy import copy
from datetime import datetime

import pandas as pd
import tkinter as tk
from tkinter import filedialog, messagebox
import openpyxl
from openpyxl.styles import Font

# ==================== 核心字段 ====================
UPDATE_FIELDS = [
    "合同号", "合同名称", "厂家", "数量", "单位", "合同签订金额",
    "付款比例", "合同签订日期", "供应商联系人及电话", "合同要求到货时间",
    "物资名称", "合同最终结算金额", "累计付款比例(百分比)", "已付金额（元）"
]

FIELD_ALIASES = {
    '合同签订金额': ['合同签订金额', '合同金额', '合同金额(元)', '合同签订金额(元)'],
}

SEGMENT_KEYWORDS = ['云南段', '广西段', '广东段']

# ==================== 工具函数 ====================
def safe_date_parse(date_str):
    if pd.isna(date_str) or str(date_str).strip() == "":
        return ""
    date_str = str(date_str).strip()
    m = re.search(r'(\d{4})[/-](\d{1,2})[/-](\d{1,2})', date_str)
    if m:
        return f"{m.group(1)}-{m.group(2).zfill(2)}-{m.group(3).zfill(2)}"
    m = re.search(r'(\d{4})年(\d{1,2})月(\d{1,2})日', date_str)
    if m:
        return f"{m.group(1)}-{m.group(2).zfill(2)}-{m.group(3).zfill(2)}"
    return date_str

def merge_materials(group):
    materials = []
    for _, row in group.iterrows():
        name = str(row.get('物资名称', '')).strip()
        qty = str(row.get('采购数量', '')).strip()
        if name and qty:
            materials.append(f"{name}({qty})")
    return "，".join(materials) if materials else ""

def aggregate_qty_unit(group):
    if group.empty or '计量单位' not in group.columns or '采购数量' not in group.columns:
        return 0, ""
    unit_sums = group.groupby('计量单位')['采购数量'].sum()
    if unit_sums.empty:
        return 0, ""
    best_unit = unit_sums.idxmax()
    best_qty = unit_sums.max()
    return best_qty, best_unit

def get_quantity_display(group):
    if group.empty or '物资名称' not in group.columns or '采购数量' not in group.columns:
        return ""
    material_groups = group.groupby('物资名称')['采购数量'].sum().reset_index()
    if material_groups.empty:
        return ""
    if len(material_groups) == 1:
        return str(material_groups.iloc[0]['采购数量'])
    lines = []
    for _, mat_row in material_groups.iterrows():
        name = str(mat_row['物资名称']).strip()
        qty = mat_row['采购数量']
        if not name:
            continue
        if ',' in name:
            name = name.split(',')[0].strip()
        elif '，' in name:
            name = name.split('，')[0].strip()
        if len(name) > 20:
            name = name[:20] + '…'
        lines.append(f"{name}（{qty}）")
    return "\n".join(lines)

def ratio_parse(ratio_str):
    if pd.isna(ratio_str):
        return ""
    ratio_str = str(ratio_str).strip()
    try:
        if ":" in ratio_str:
            parts = [p.strip() for p in ratio_str.split(":") if p.strip()]
            if parts:
                total = sum(float(x) for x in parts)
                return "/".join([f"{float(x)/total*100:.1f}%" for x in parts])
        if "%" in ratio_str:
            return ratio_str
        ratio = float(ratio_str)
        if 0 <= ratio <= 1:
            return f"{ratio*100:.1f}%"
        return f"{ratio}%"
    except:
        return ratio_str

def percent_to_decimal(pct_str):
    if not pct_str or pd.isna(pct_str):
        return 0.0
    pct_str = str(pct_str).strip().replace('%', '')
    try:
        return float(pct_str) / 100.0
    except:
        return 0.0

def values_differ(old_val, new_val):
    if old_val is None:
        old_val = ""
    if new_val is None:
        new_val = ""
    try:
        return float(old_val) != float(new_val)
    except (ValueError, TypeError):
        pass
    def to_str(v):
        if isinstance(v, datetime):
            return v.strftime("%Y-%m-%d")
        return str(v).strip()
    return to_str(old_val) != to_str(new_val)

def copy_row_style(ws, from_row, to_row, max_col=None):
    if max_col is None:
        max_col = ws.max_column
    for col in range(1, max_col+1):
        src = ws.cell(row=from_row, column=col)
        dst = ws.cell(row=to_row, column=col)
        dst.font = copy(src.font)
        dst.border = copy(src.border)
        dst.fill = copy(src.fill)
        dst.alignment = copy(src.alignment)
        dst.number_format = src.number_format

def get_first_column_as_sequence(ws, data_start, max_row):
    header_cell = ws.cell(row=data_start-1, column=1)
    if header_cell.value:
        clean_col = re.sub(r'[()（）\s]', '', str(header_cell.value)).lower()
        for field in UPDATE_FIELDS:
            if re.sub(r'[()（）\s]', '', field).lower() == clean_col:
                return None
    if max_row >= data_start:
        cell = ws.cell(row=max_row, column=1)
        try:
            return int(cell.value) if cell.value is not None else 0
        except:
            return None
    return 0

def is_file_writable(filepath):
    try:
        with open(filepath, 'a'):
            pass
        return True
    except (IOError, PermissionError):
        return False

def read_engineering_list(filepath):
    """读取工程列表（第一行为列名），返回项目名称列表"""
    df = pd.read_excel(filepath)
    df.columns = df.columns.astype(str).str.strip()
    name_col = None
    for col in df.columns:
        clean = re.sub(r'[()（）\s]', '', col).lower()
        if '工程项目名称' in clean or '项目名称' in clean:
            name_col = col
            break
    if not name_col:
        raise ValueError("工程列表文件中未找到“工程项目名称”或“项目名称”列")
    projects = df[name_col].dropna().astype(str).str.strip().tolist()
    # 去重保序
    seen = set()
    unique = []
    for p in projects:
        if p not in seen:
            seen.add(p)
            unique.append(p)
    return unique

def clean_project_name(name):
    """去掉项目名称中括号及其内容，用于模糊匹配"""
    return re.sub(r'[（(].*?[）)]', '', str(name)).strip()


def find_all_summary_files(summary_folder, project_name):
    """在文件夹中查找所有包含项目名称的xlsx文件（支持多模板）"""
    clean_name = clean_project_name(project_name)
    result = []
    for fname in os.listdir(summary_folder):
        if fname.endswith('.xlsx') and clean_name in fname and not fname.startswith('~$'):
            result.append(os.path.join(summary_folder, fname))
    return result


def _has_segment_sheets(filepath):
    """检查模板文件是否包含分段sheet（线路材料模板特征）"""
    try:
        wb = openpyxl.load_workbook(filepath)
        for sn in wb.sheetnames:
            for seg in SEGMENT_KEYWORDS:
                if seg in sn:
                    wb.close()
                    return True
        wb.close()
    except Exception:
        pass
    return False


def classify_templates(file_paths):
    """将模板文件分类为线路材料模板和设备模板，返回 (line_material_path, device_path)"""
    line_path = None
    device_path = None
    for fp in file_paths:
        if _has_segment_sheets(fp):
            line_path = fp
        else:
            device_path = fp
    return line_path, device_path


def extract_segment(name):
    """从单项工程名称中提取段名，如 线路工程（云南段） -> 云南段"""
    name_str = str(name)
    for seg in SEGMENT_KEYWORDS:
        if seg in name_str:
            return seg
    return None


def extract_template_tag(filepath):
    """从模板文件名中提取 [xxx] 标签"""
    fname = os.path.basename(filepath)
    m = re.search(r'\[.*?\]', fname)
    return m.group(0) if m else ''


def split_data_by_sub_project(main_proj, sub_proj):
    """按单项工程名称是否包含线路工程拆分主表和次表数据"""
    main_sub_col = None
    for col in main_proj.columns:
        if '单项工程名称' in str(col):
            main_sub_col = col
            break
    sub_sub_col = None
    for col in sub_proj.columns:
        if '单项工程名称' in str(col):
            sub_sub_col = col
            break

    if main_sub_col is None or sub_sub_col is None:
        return (main_proj, sub_proj), (None, None)

    main_line = main_proj[main_proj[main_sub_col].astype(str).str.contains('线路工程')].copy()
    main_device = main_proj[~main_proj[main_sub_col].astype(str).str.contains('线路工程')].copy()

    line_contracts = set(main_line['合同编号'].astype(str).str.strip()) if '合同编号' in main_line.columns else set()
    device_contracts = set(main_device['合同编号'].astype(str).str.strip()) if '合同编号' in main_device.columns else set()

    sub_line = pd.DataFrame()
    sub_device = pd.DataFrame()
    if '合同编号' in sub_proj.columns:
        sub_proj_copy = sub_proj.copy()
        sub_proj_copy['_contract'] = sub_proj_copy['合同编号'].astype(str).str.strip()
        sub_line = sub_proj_copy[sub_proj_copy['_contract'].isin(line_contracts)].drop(columns=['_contract'])
        sub_device = sub_proj_copy[sub_proj_copy['_contract'].isin(device_contracts)].drop(columns=['_contract'])

    return (main_line, sub_line), (main_device, sub_device)


def clean_filename(name):
    """将字符串转换为合法的文件名（保留中文、字母、数字、空格、-、_）"""
    name = re.sub(r'[\\/*?:"<>|]', '_', name)
    name = name.strip('. ')
    if not name:
        name = "未命名"
    return name

# ==================== 构建结果 DataFrame ====================
def _build_result(main_proj, sub_proj):
    """合并主次表，生成结果DataFrame"""
    sub_grouped = sub_proj.groupby('合同编号')
    sub_agg = []
    for contract, group in sub_grouped:
        qty, unit = aggregate_qty_unit(group)
        mats = merge_materials(group)
        qty_disp = get_quantity_display(group)
        total = group['价税合计(元)'].sum() if '价税合计(元)' in group.columns else 0
        dates = group['合同最终交货日期'].dropna()
        deliv = dates.iloc[0] if not dates.empty else ""
        sub_agg.append({
            '合同编号': contract,
            '物资名称': mats,
            '数量': qty_disp,
            '单位': unit,
            '合同最终结算金额': total,
            '合同最终交货日期': deliv
        })
    sub_processed = pd.DataFrame(sub_agg)

    merged = pd.merge(main_proj, sub_processed, on='合同编号', how='left')

    result = pd.DataFrame()
    result['合同号'] = merged['合同编号']
    result['合同名称'] = merged['合同名称']
    result['厂家'] = merged['供应商名称']
    result['数量'] = merged['数量']
    result['单位'] = merged['单位']
    result['合同签订金额'] = merged['合同金额(元)']
    result['合同签订日期'] = merged['合同签约完成时间']
    contact = merged.get('供应商联系人', pd.Series('')).fillna('').astype(str)
    phone = merged.get('供应商联系人电话', pd.Series('')).fillna('').astype(str)
    result['供应商联系人及电话'] = contact + phone
    result['合同要求到货时间'] = merged['合同最终交货日期'].apply(safe_date_parse)
    result['付款比例'] = merged['合同支付比例']
    result['物资名称'] = merged['物资名称']
    result['合同最终结算金额'] = merged['合同最终结算金额']
    result['累计付款比例(百分比)'] = merged['付款比例']

    # 投标偏差
    remarks = merged.get('备注', pd.Series('')).fillna('').astype(str)
    def get_bid_deviation(remark):
        if '正偏差' in remark:
            return '正偏差'
        elif '负偏差' in remark:
            return '负偏差'
        else:
            return '否'
    result['投标偏差（是/否）'] = remarks.apply(get_bid_deviation)

    try:
        result['_sort'] = pd.to_datetime(result['合同签订日期'], errors='coerce')
        result = result.sort_values('_sort', na_position='last').drop(columns=['_sort'])
    except:
        result = result.sort_values('合同签订日期')
    result = result.reset_index(drop=True)
    return result


# ==================== 模板更新 ====================
def _update_template(result, summary_path, project_name):
    """更新单个汇总表模板（自动判断单sheet/多sheet）"""
    if not os.path.exists(summary_path):
        return False, f"汇总表文件不存在: {summary_path}"
    wb = openpyxl.load_workbook(summary_path)
    segment_sheets = {}
    for sn in wb.sheetnames:
        for seg in SEGMENT_KEYWORDS:
            if seg in sn:
                segment_sheets[seg] = sn
                break
    if len(segment_sheets) >= 2:
        return _process_multi_sheet(wb, result, segment_sheets, summary_path, project_name)
    else:
        return _process_single_sheet(wb, result, summary_path, project_name)


def _process_single_sheet(wb, result, summary_path, project_name):
    """单sheet模板：更新活动sheet"""
    ws = wb.active
    header_row = None
    for row in ws.iter_rows(min_row=1, max_row=10):
        for cell in row:
            if cell.value and "合同号" in str(cell.value):
                header_row = cell.row
                break
        if header_row:
            break
    if header_row is None:
        header_row = 4

    col_map = {}
    for cell in ws[header_row]:
        if cell.value:
            col_name = str(cell.value).strip()
            clean_col = re.sub(r'[()（）\s]', '', col_name).lower()
            for field in UPDATE_FIELDS:
                if field in col_map:
                    continue
                aliases = FIELD_ALIASES.get(field, [field])
                for alias in aliases:
                    if re.sub(r'[()（）\s]', '', alias).lower() == clean_col:
                        col_map[field] = cell.column
                        break
    if '合同号' not in col_map:
        return False, "汇总表中未找到合同号列"

    data_start = header_row + 1
    contract_row_map = {}
    for row in ws.iter_rows(min_row=data_start, max_row=ws.max_row,
                            min_col=col_map['合同号'], max_col=col_map['合同号']):
        cell = row[0]
        if cell.value is not None:
            contract_row_map[str(cell.value).strip()] = cell.row

    blue = "0000FF"
    green = "008000"
    yellow = "FFA500"
    updated_count = 0
    changed_contracts = []
    new_rows_data = []

    for _, r_row in result.iterrows():
        contract = str(r_row['合同号']).strip()
        if contract in contract_row_map:
            row_idx = contract_row_map[contract]
            row_changed = False
            for field in UPDATE_FIELDS:
                if field == '已付金额（元）':
                    continue
                if field == '供应商联系人及电话':
                    continue
                if field in col_map and field in r_row:
                    col_idx = col_map[field]
                    old_cell = ws.cell(row=row_idx, column=col_idx)
                    if values_differ(old_cell.value, r_row[field]):
                        old_cell.value = r_row[field]
                        new_font = copy(old_cell.font) if old_cell.font else Font()
                        new_font.color = blue
                        old_cell.font = new_font
                        row_changed = True
            if '已付金额（元）' in col_map:
                amt_col = col_map.get('合同签订金额')
                rat_col = col_map.get('累计付款比例(百分比)')
                if amt_col and rat_col:
                    amt_val = ws.cell(row=row_idx, column=amt_col).value
                    rat_val = ws.cell(row=row_idx, column=rat_col).value
                    try:
                        amt = float(amt_val) if amt_val not in [None, ""] else 0
                        rat = percent_to_decimal(rat_val)
                        paid = amt * rat
                        paid_cell = ws.cell(row=row_idx, column=col_map['已付金额（元）'])
                        if values_differ(paid_cell.value, paid):
                            paid_cell.value = paid
                            paid_cell.number_format = '#,##0.00'
                            new_font = copy(paid_cell.font) if paid_cell.font else Font()
                            new_font.color = blue
                            paid_cell.font = new_font
                            row_changed = True
                    except:
                        pass
            if row_changed:
                updated_count += 1
                changed_contracts.append(contract)
        else:
            new_rows_data.append(r_row)

    if new_rows_data:
        template_row = ws.max_row if ws.max_row >= data_start else data_start
        seq_index = get_first_column_as_sequence(ws, data_start, ws.max_row)
        seq = seq_index if seq_index is not None else 0
        for r_row in new_rows_data:
            ins_row = ws.max_row + 1
            copy_row_style(ws, template_row, ins_row)
            if seq_index is not None:
                seq += 1
                seq_cell = ws.cell(row=ins_row, column=1)
                seq_cell.value = seq
                seq_cell.font = Font(color=green)
            for field in UPDATE_FIELDS:
                if field == '已付金额（元）':
                    continue
                if field in col_map and field in r_row:
                    cell = ws.cell(row=ins_row, column=col_map[field])
                    cell.value = r_row[field]
                    cell.font = Font(color=green)
            if '已付金额（元）' in col_map:
                amt_col = col_map.get('合同签订金额')
                rat_col = col_map.get('累计付款比例(百分比)')
                if amt_col and rat_col:
                    amt = float(r_row['合同签订金额']) if r_row['合同签订金额'] not in [None, ""] else 0
                    rat = percent_to_decimal(r_row['累计付款比例(百分比)'])
                    paid = amt * rat
                    paid_cell = ws.cell(row=ins_row, column=col_map['已付金额（元）'])
                    paid_cell.value = paid
                    paid_cell.number_format = '#,##0.00'
                    paid_cell.font = Font(color=green)
            template_row = ins_row

    result_contracts = set(result['合同号'].astype(str).str.strip())
    yellow_font = Font(color=yellow)
    for row in ws.iter_rows(min_row=data_start, max_row=ws.max_row,
                            min_col=col_map['合同号'], max_col=col_map['合同号']):
        cell = row[0]
        if cell.value is not None:
            c = str(cell.value).strip()
            if c and c not in result_contracts:
                for col_idx in range(1, ws.max_column + 1):
                    target = ws.cell(row=cell.row, column=col_idx)
                    if target.value is not None:
                        new_font = copy(target.font) if target.font else Font()
                        new_font.color = yellow
                        target.font = new_font

    wb.save(summary_path)

    changed_info = "、".join(changed_contracts) if changed_contracts else "无"
    return True, (f"项目 {project_name} 更新完成\n"
                  f"更新合同: {updated_count} 个\n"
                  f"新增合同: {len(new_rows_data)} 个\n"
                  f"变更合同号: {changed_info}")


def _process_multi_sheet(wb, result, segment_sheets, summary_path, project_name):
    """多sheet分段模板：遍历分段sheet，匹配合同号并更新"""
    blue = "0000FF"
    green = "008000"
    yellow = "FFA500"

    sheet_info = {}
    for seg, sn in segment_sheets.items():
        ws = wb[sn]
        header_row = None
        for row in ws.iter_rows(min_row=1, max_row=10):
            for cell in row:
                if cell.value and "合同号" in str(cell.value):
                    header_row = cell.row
                    break
            if header_row:
                break
        if header_row is None:
            header_row = 4

        col_map = {}
        for cell in ws[header_row]:
            if cell.value:
                col_name = str(cell.value).strip()
                clean_col = re.sub(r'[()（）\s]', '', col_name).lower()
                for field in UPDATE_FIELDS:
                    if field in col_map:
                        continue
                    aliases = FIELD_ALIASES.get(field, [field])
                    for alias in aliases:
                        if re.sub(r'[()（）\s]', '', alias).lower() == clean_col:
                            col_map[field] = cell.column
                            break

        data_start = header_row + 1
        contract_row_map = {}
        if '合同号' in col_map:
            for row in ws.iter_rows(min_row=data_start, max_row=ws.max_row,
                                    min_col=col_map['合同号'], max_col=col_map['合同号']):
                cell = row[0]
                if cell.value is not None:
                    contract_row_map[str(cell.value).strip()] = cell.row

        sheet_info[seg] = {
            'ws': ws, 'col_map': col_map, 'contract_row_map': contract_row_map,
            'header_row': header_row, 'data_start': data_start
        }

    global_contract_map = {}
    for seg, info in sheet_info.items():
        for contract in info['contract_row_map']:
            if contract not in global_contract_map:
                global_contract_map[contract] = seg

    def determine_segment(contract_name):
        name = str(contract_name)
        for seg in SEGMENT_KEYWORDS:
            if seg in name:
                return seg
        return None

    updated_count = 0
    new_count = 0
    changed_contracts = []
    new_contracts_data = {}

    for _, r_row in result.iterrows():
        contract = str(r_row['合同号']).strip()
        if contract in global_contract_map:
            seg = global_contract_map[contract]
            info = sheet_info[seg]
            ws = info['ws']
            col_map = info['col_map']
            row_idx = info['contract_row_map'][contract]
            row_changed = False
            for field in UPDATE_FIELDS:
                if field == '已付金额（元）':
                    continue
                if field == '供应商联系人及电话':
                    continue
                if field in col_map and field in r_row:
                    col_idx = col_map[field]
                    old_cell = ws.cell(row=row_idx, column=col_idx)
                    if values_differ(old_cell.value, r_row[field]):
                        old_cell.value = r_row[field]
                        new_font = copy(old_cell.font) if old_cell.font else Font()
                        new_font.color = blue
                        old_cell.font = new_font
                        row_changed = True
            if '已付金额（元）' in col_map:
                amt_col = col_map.get('合同签订金额')
                rat_col = col_map.get('累计付款比例(百分比)')
                if amt_col and rat_col:
                    amt_val = ws.cell(row=row_idx, column=amt_col).value
                    rat_val = ws.cell(row=row_idx, column=rat_col).value
                    try:
                        amt = float(amt_val) if amt_val not in [None, ""] else 0
                        rat = percent_to_decimal(rat_val)
                        paid = amt * rat
                        paid_cell = ws.cell(row=row_idx, column=col_map['已付金额（元）'])
                        if values_differ(paid_cell.value, paid):
                            paid_cell.value = paid
                            paid_cell.number_format = '#,##0.00'
                            new_font = copy(paid_cell.font) if paid_cell.font else Font()
                            new_font.color = blue
                            paid_cell.font = new_font
                            row_changed = True
                    except:
                        pass
            if row_changed:
                updated_count += 1
                changed_contracts.append(contract)
        else:
            target_seg = determine_segment(r_row['合同名称'])
            if target_seg and target_seg in sheet_info:
                if target_seg not in new_contracts_data:
                    new_contracts_data[target_seg] = []
                new_contracts_data[target_seg].append(r_row)
                new_count += 1

    for seg, rows in new_contracts_data.items():
        info = sheet_info[seg]
        ws = info['ws']
        col_map = info['col_map']
        data_start = info['data_start']
        if '合同号' not in col_map:
            continue
        template_row = ws.max_row if ws.max_row >= data_start else data_start
        seq_index = get_first_column_as_sequence(ws, data_start, ws.max_row)
        seq = seq_index if seq_index is not None else 0
        for r_row in rows:
            ins_row = ws.max_row + 1
            copy_row_style(ws, template_row, ins_row)
            if seq_index is not None:
                seq += 1
                seq_cell = ws.cell(row=ins_row, column=1)
                seq_cell.value = seq
                seq_cell.font = Font(color=green)
            for field in UPDATE_FIELDS:
                if field == '已付金额（元）':
                    continue
                if field in col_map and field in r_row:
                    cell = ws.cell(row=ins_row, column=col_map[field])
                    cell.value = r_row[field]
                    cell.font = Font(color=green)
            if '已付金额（元）' in col_map:
                amt_col = col_map.get('合同签订金额')
                rat_col = col_map.get('累计付款比例(百分比)')
                if amt_col and rat_col:
                    amt = float(r_row['合同签订金额']) if r_row['合同签订金额'] not in [None, ""] else 0
                    rat = percent_to_decimal(r_row['累计付款比例(百分比)'])
                    paid = amt * rat
                    paid_cell = ws.cell(row=ins_row, column=col_map['已付金额（元）'])
                    paid_cell.value = paid
                    paid_cell.number_format = '#,##0.00'
                    paid_cell.font = Font(color=green)
            template_row = ins_row

    result_contracts = set(result['合同号'].astype(str).str.strip())
    yellow_font = Font(color=yellow)
    for seg, info in sheet_info.items():
        ws = info['ws']
        col_map = info['col_map']
        data_start = info['data_start']
        if '合同号' not in col_map:
            continue
        for row in ws.iter_rows(min_row=data_start, max_row=ws.max_row,
                                min_col=col_map['合同号'], max_col=col_map['合同号']):
            cell = row[0]
            if cell.value is not None:
                c = str(cell.value).strip()
                if c and c not in result_contracts:
                    for col_idx in range(1, ws.max_column + 1):
                        target = ws.cell(row=cell.row, column=col_idx)
                        if target.value is not None:
                            new_font = copy(target.font) if target.font else Font()
                            new_font.color = yellow
                            target.font = new_font

    wb.save(summary_path)

    changed_info = "、".join(changed_contracts) if changed_contracts else "无"
    return True, (f"项目 {project_name} 更新完成\n"
                  f"更新合同: {updated_count} 个\n"
                  f"新增合同: {new_count} 个\n"
                  f"变更合同号: {changed_info}")


# ==================== 单项目处理（修改新文件） ====================
def process_project(main_df, sub_df, project_name, summary_path, device_template_path=None):
    """
    直接在 summary_path 上修改，不再创建备份。
    若提供 device_template_path，则按单项工程名称拆分数据，
    分别写入线路材料模板和设备模板。
    """
    try:
        # ---- 1. 筛选主表和次表 ----
        main_name_col = None
        for col in main_df.columns:
            clean = re.sub(r'[()（）\s]', '', str(col)).lower()
            if '项目名称' in clean or '工程名称' in clean:
                main_name_col = col
                break
        if not main_name_col:
            return False, "主表中未找到项目名称列"
        sub_name_col = None
        for col in sub_df.columns:
            clean = re.sub(r'[()（）\s]', '', str(col)).lower()
            if '项目名称' in clean or '工程名称' in clean:
                sub_name_col = col
                break
        if not sub_name_col:
            return False, "次表中未找到项目名称列"

        proj_str = str(project_name).strip()
        clean_proj_name = clean_project_name(proj_str)
        main_mask = main_df[main_name_col].astype(str).str.strip().apply(
            lambda x: x == proj_str or x == clean_proj_name)
        sub_mask = sub_df[sub_name_col].astype(str).str.strip().apply(
            lambda x: x == proj_str or x == clean_proj_name)
        main_proj = main_df[main_mask].copy()
        sub_proj = sub_df[sub_mask].copy()
        if main_proj.empty and sub_proj.empty:
            return False, f"项目 {project_name} 在主表和次表中均无数据"

        # ---- 2. 主表列名标准化 ----
        if not main_proj.empty:
            required_main = {
                '合同编号': ['合同编号'],
                '合同名称': ['合同名称'],
                '供应商名称': ['供应商名称'],
                '合同金额(元)': ['合同金额(元)', '合同金额', '金额(元)'],
                '合同签约完成时间': ['合同签约完成时间', '签约完成时间', '合同签订日期'],
                '合同支付比例': ['合同支付比例', '支付比例'],
                '付款比例': ['付款比例'],
                '供应商联系人': ['供应商联系人', '联系人'],
                '供应商联系人电话': ['供应商联系人电话', '联系人电话'],
                '备注': ['备注'],
            }
            for std, aliases in required_main.items():
                if std in main_proj.columns:
                    continue
                found = False
                for alias in aliases:
                    if alias in main_proj.columns:
                        main_proj.rename(columns={alias: std}, inplace=True)
                        found = True
                        break
                if not found:
                    clean_std = re.sub(r'[()（）\s]', '', std).lower()
                    for act in main_proj.columns:
                        if re.sub(r'[()（）\s]', '', str(act)).lower() == clean_std:
                            main_proj.rename(columns={act: std}, inplace=True)
                            found = True
                            break
                if not found and std in ['合同编号','合同名称','供应商名称','合同金额(元)','合同签约完成时间']:
                    return False, f"主表缺少必要列: {std}\n实际列名:\n" + "\n".join(main_proj.columns)

        if not sub_proj.empty:
            sub_required = ['合同编号', '物资名称', '采购数量', '计量单位', '价税合计(元)', '合同最终交货日期']
            for col in sub_required:
                if col not in sub_proj.columns:
                    matched = False
                    clean_need = re.sub(r'[()（）\s]', '', col).lower()
                    for act in sub_proj.columns:
                        if re.sub(r'[()（）\s]', '', str(act)).lower() == clean_need:
                            sub_proj.rename(columns={act: col}, inplace=True)
                            matched = True
                            break
                    if not matched:
                        return False, f"次表缺少必要列: {col}"

        main_proj['合同编号'] = main_proj['合同编号'].astype(str).str.strip()
        sub_proj['合同编号'] = sub_proj['合同编号'].astype(str).str.strip()

        # ---- 3. 判断是否多模板处理 ----
        if device_template_path and os.path.exists(device_template_path):
            (main_line, sub_line), (main_device, sub_device) = split_data_by_sub_project(main_proj, sub_proj)
            msgs = []
            all_ok = True

            if not main_line.empty and not sub_line.empty:
                result = _build_result(main_line, sub_line)
                ok, msg = _update_template(result, summary_path, project_name)
                all_ok = all_ok and ok
                msgs.append(f"[线路材料] {msg}")
            else:
                msgs.append("[线路材料] 无线路工程数据，跳过")

            if not main_device.empty and not sub_device.empty:
                result = _build_result(main_device, sub_device)
                ok, msg = _update_template(result, device_template_path, project_name)
                all_ok = all_ok and ok
                msgs.append(f"[设备] {msg}")
            else:
                msgs.append("[设备] 无设备数据，跳过")

            return all_ok, "\n".join(msgs)
        else:
            if main_proj.empty or sub_proj.empty:
                return False, f"项目 {project_name} 在主表或次表中无数据"
            result = _build_result(main_proj, sub_proj)
            return _update_template(result, summary_path, project_name)
    except Exception as e:
        return False, f"处理 {project_name} 失败:\n{str(e)}\n\n{traceback.format_exc()}"

# ==================== GUI ====================
class App:
    def __init__(self, root):
        self.root = root
        self.root.title("合同数据合并工具")
        self.root.geometry("650x750")
        self.root.configure(bg="#f0f0f0")

        self.main_var = tk.StringVar()
        self.sub_var = tk.StringVar()
        self.eng_var = tk.StringVar()
        self.folder_var = tk.StringVar()
        self.time_var = tk.StringVar(value=datetime.now().strftime("%Y%m%d"))

        self.project_names = []

        tk.Label(root, text="合同数据合并工具", font=("微软雅黑", 16, "bold"),
                 bg="#f0f0f0", fg="#2c3e50").pack(pady=15)

        frame = tk.LabelFrame(root, text="文件选择", bg="#f0f0f0", padx=10, pady=10)
        frame.pack(fill="x", padx=20, pady=5)

        self._file_row(frame, "工程列表:", self.eng_var, self.load_engineering)
        self._file_row(frame, "采购合同明细:", self.main_var)
        self._file_row(frame, "采购合同查询明细:", self.sub_var)
        self._folder_row(frame, "汇总表文件夹:", self.folder_var)

        # 自定义后缀
        time_frame = tk.Frame(frame, bg="#f0f0f0")
        time_frame.pack(fill="x", pady=5)
        tk.Label(time_frame, text="自定义后缀:", bg="#f0f0f0", width=18, anchor="w").pack(side="left")
        tk.Entry(time_frame, textvariable=self.time_var, width=20).pack(side="left", padx=5)
        tk.Label(time_frame, text="（用于输出文件名，例如 20260521）",
                 bg="#f0f0f0", fg="#555555", font=("微软雅黑", 9)).pack(side="left", padx=5)

        proj_frame = tk.LabelFrame(root, text="选择工程项目（多选）", bg="#f0f0f0", padx=10, pady=10)
        proj_frame.pack(fill="both", expand=True, padx=20, pady=5)
        list_frame = tk.Frame(proj_frame)
        list_frame.pack(fill="both", expand=True)
        scroll = tk.Scrollbar(list_frame)
        scroll.pack(side="right", fill="y")
        self.listbox = tk.Listbox(list_frame, selectmode=tk.MULTIPLE, yscrollcommand=scroll.set,
                                  font=("微软雅黑", 10), width=40, height=8)
        self.listbox.pack(side="left", fill="both", expand=True)
        scroll.config(command=self.listbox.yview)
        btnf = tk.Frame(proj_frame, bg="#f0f0f0")
        btnf.pack(pady=5)
        tk.Button(btnf, text="全选", command=lambda: self.listbox.select_set(0, tk.END)).pack(side="left", padx=5)
        tk.Button(btnf, text="取消全选", command=lambda: self.listbox.selection_clear(0, tk.END)).pack(side="left", padx=5)

        self.run_btn = tk.Button(root, text="开始更新汇总表", command=self.run,
                                 bg="#4CAF50", fg="white", font=("微软雅黑", 12, "bold"),
                                 width=20, height=2)
        self.run_btn.pack(pady=15)

        self.status = tk.StringVar()
        self.status.set("请加载工程列表并选择项目")
        tk.Label(root, textvariable=self.status, bg="#f0f0f0", fg="#555555").pack(pady=5)

    def _file_row(self, parent, label, var, cmd=None):
        f = tk.Frame(parent, bg="#f0f0f0")
        f.pack(fill="x", pady=5)
        tk.Label(f, text=label, bg="#f0f0f0", width=18, anchor="w").pack(side="left")
        e = tk.Entry(f, textvariable=var, width=40)
        e.pack(side="left", padx=5)
        tk.Button(f, text="浏览...", command=cmd if cmd else lambda: self.browse_file(var)).pack(side="left")

    def _folder_row(self, parent, label, var):
        f = tk.Frame(parent, bg="#f0f0f0")
        f.pack(fill="x", pady=5)
        tk.Label(f, text=label, bg="#f0f0f0", width=18, anchor="w").pack(side="left")
        e = tk.Entry(f, textvariable=var, width=40)
        e.pack(side="left", padx=5)
        tk.Button(f, text="浏览...", command=lambda: self.browse_folder(var)).pack(side="left")

    def browse_file(self, var):
        path = filedialog.askopenfilename(filetypes=[("Excel文件", "*.xlsx;*.xls")])
        if path:
            var.set(path)

    def browse_folder(self, var):
        path = filedialog.askdirectory()
        if path:
            var.set(path)

    def load_engineering(self):
        path = filedialog.askopenfilename(filetypes=[("Excel文件", "*.xlsx;*.xls")])
        if not path:
            return
        self.eng_var.set(path)
        try:
            self.project_names = read_engineering_list(path)
            self.listbox.delete(0, tk.END)
            for name in self.project_names:
                self.listbox.insert(tk.END, name)
            self.status.set(f"已加载 {len(self.project_names)} 个项目")
        except Exception as e:
            messagebox.showerror("错误", f"加载工程列表失败:\n{str(e)}")

    def run(self):
        main = self.main_var.get()
        sub = self.sub_var.get()
        eng = self.eng_var.get()
        folder = self.folder_var.get()
        time_suffix = self.time_var.get().strip()

        if not all([main, sub, eng, folder]):
            messagebox.showerror("错误", "请选择所有文件/文件夹")
            return
        if not time_suffix:
            messagebox.showerror("错误", "请填写自定义后缀（例如 20260521）")
            return

        selected = self.listbox.curselection()
        if not selected:
            messagebox.showerror("错误", "请至少选择一个项目")
            return
        projects = [self.project_names[i] for i in selected]

        clean_time = clean_filename(time_suffix)
        if not clean_time:
            clean_time = datetime.now().strftime("%Y%m%d")

        # 创建输出文件夹
        output_dir = os.path.join(folder, "更新结果")
        os.makedirs(output_dir, exist_ok=True)

        self.run_btn.config(state=tk.DISABLED, text="处理中...")
        self.root.update()

        try:
            main_df = pd.read_excel(main, skiprows=1)
            sub_df = pd.read_excel(sub, skiprows=1)
            for df in (main_df, sub_df):
                df.columns = df.columns.str.strip()
                df.columns = df.columns.str.replace('（', '(').str.replace('）', ')')
                df.columns = df.columns.str.replace('\u200b', '')

            success = []
            fail = []
            for proj in projects:
                template_files = find_all_summary_files(folder, proj)
                if not template_files:
                    fail.append(f"{proj}: 未找到汇总表文件")
                    continue

                clean_proj = clean_filename(proj)

                if len(template_files) >= 2:
                    line_path, device_path = classify_templates(template_files)
                    if not line_path or not device_path:
                        fail.append(f"{proj}: 检测到{len(template_files)}个模板，无法区分")
                        continue

                    line_tag = extract_template_tag(line_path)
                    device_tag = extract_template_tag(device_path)

                    line_new_path = os.path.join(output_dir, f"{clean_proj}{line_tag}物资管理一本账{clean_time}.xlsx")
                    device_new_path = os.path.join(output_dir, f"{clean_proj}{device_tag}物资管理一本账{clean_time}.xlsx")

                    shutil.copyfile(line_path, line_new_path)
                    shutil.copyfile(device_path, device_new_path)

                    ok, msg = process_project(main_df.copy(), sub_df.copy(), proj,
                                               line_new_path, device_new_path)
                    if ok:
                        success.append(msg)
                    else:
                        for p in [line_new_path, device_new_path]:
                            if os.path.exists(p):
                                os.remove(p)
                        fail.append(f"{proj}: {msg}")
                else:
                    original_path = template_files[0]
                    new_filename = f"{clean_proj}物资管理一本账{clean_time}.xlsx"
                    new_path = os.path.join(output_dir, new_filename)
                    shutil.copyfile(original_path, new_path)
                    ok, msg = process_project(main_df.copy(), sub_df.copy(), proj, new_path)
                    if ok:
                        success.append(msg)
                    else:
                        if os.path.exists(new_path):
                            os.remove(new_path)
                        fail.append(f"{proj}: {msg}")

            result_text = f"成功 {len(success)} 个项目"
            if fail:
                result_text += "\n\n失败项目:\n" + "\n".join(fail)
            messagebox.showinfo("处理结果", result_text)
            self.status.set("处理完成")
        except Exception as e:
            messagebox.showerror("错误", f"运行失败:\n{str(e)}")
        finally:
            self.run_btn.config(state=tk.NORMAL, text="开始更新汇总表")


if __name__ == "__main__":
    root = tk.Tk()
    app = App(root)
    root.mainloop()